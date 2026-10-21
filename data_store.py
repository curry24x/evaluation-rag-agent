from __future__ import annotations

import contextvars
import hashlib
import json
import re
import time
from functools import lru_cache
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from fastembed import TextEmbedding
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import load_workbook

PROJECT_DIR = Path(__file__).resolve().parent
DATA_DIR = PROJECT_DIR / "data"
DB_DIR = PROJECT_DIR / ".chroma"
MODEL_DIR = PROJECT_DIR / ".models"
EMBEDDING_MODEL = "BAAI/bge-small-zh-v1.5"
load_dotenv(PROJECT_DIR / ".env")

_DATA_FILES = ("200角色看图打分工作簿.xlsx", "error_analysis_summary.md",
               "error_analysis_per_image.csv", "error_pattern_summary.csv",
               "qwen_doubao_task_accuracy.csv", "task_model_accuracy_4x4.csv")
_DOC_CACHE: list[Document] | None = None
_DOC_CACHE_SIGNATURE: tuple[Any, ...] | None = None
_REQUEST_METRICS: contextvars.ContextVar[dict[str, float] | None] = contextvars.ContextVar(
    "request_metrics", default=None)


def add_timing(name: str, milliseconds: float) -> None:
    metrics = _REQUEST_METRICS.get()
    if metrics is not None:
        metrics[name] = metrics.get(name, 0.0) + milliseconds


class ChineseSemanticEmbeddings(Embeddings):
    """FastEmbed wrapper around the local pretrained Chinese BGE model."""

    def __init__(self):
        started = time.perf_counter()
        self.model = TextEmbedding(model_name=EMBEDDING_MODEL, cache_dir=str(MODEL_DIR))
        add_timing("embedding_ms", (time.perf_counter() - started) * 1000)

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        started = time.perf_counter()
        result = [vector.tolist() for vector in self.model.passage_embed(texts)]
        add_timing("embedding_ms", (time.perf_counter() - started) * 1000)
        return result

    def embed_query(self, text: str) -> list[float]:
        started = time.perf_counter()
        result = next(self.model.query_embed([text])).tolist()
        add_timing("embedding_ms", (time.perf_counter() - started) * 1000)
        return result


def data_signature() -> tuple[Any, ...]:
    signature = []
    for filename in _DATA_FILES:
        path = DATA_DIR / filename
        try:
            info = path.stat()
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            signature.append((filename, info.st_size, info.st_mtime_ns, digest))
        except FileNotFoundError:
            signature.append((filename, None, None, None))
    return tuple(signature)


def read_csv(name: str) -> list[dict[str, str]]:
    import csv
    with (DATA_DIR / name).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


_read_csv = read_csv


def _workbook_documents() -> list[Document]:
    path = DATA_DIR / "200角色看图打分工作簿.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    documents: list[Document] = []
    rules = workbook["判读规则"]
    for row in rules.iter_rows(min_row=4, max_row=7, values_only=True):
        tier, index_range, question, redline, example = row[:5]
        content = (f"{tier} 评测规则\n序号范围：{index_range}\n核心问题：{question}\n"
                   f"判错红线：{redline}\n示例：{example}")
        documents.append(Document(content, metadata={"source": path.name, "section": str(tier),
            "document_type": "tier_rule", "tier": str(tier), "task": f"Task {str(tier).split()[-1]}"}))
    operation_rows = [str(row[0]) for row in rules.iter_rows(min_row=10, values_only=True) if row[0]]
    documents.append(Document("人工评测操作规则\n" + "\n".join(operation_rows),
        metadata={"source": path.name, "section": "操作规则", "document_type": "operation_rule"}))
    checklist = workbook["看图清单"]
    for row in checklist.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        index, tier, character, category, atomic, redline, note = row[:7]
        content = (f"样本 #{index}｜{tier}｜{character}\n作品/类别：{category}\n"
                   f"核对顺序（原子特征）：\n{atomic}\n怎么看（判错红线）：\n{redline}")
        if note:
            content += f"\n备注：{note}"
        metadata = {"source": path.name, "section": f"样本{index}", "document_type": "checklist",
                    "tier": str(tier), "task": f"Task {str(tier).split()[-1]}", "index": int(index)}
        if character:
            metadata["character"] = str(character)
        if category:
            metadata["category"] = str(category)
        documents.append(Document(content, metadata=metadata))
    workbook.close()
    return documents


def _markdown_documents() -> list[Document]:
    path = DATA_DIR / "error_analysis_summary.md"
    text = path.read_text(encoding="utf-8-sig")
    sections = re.split(r"(?=^#{1,3} )", text, flags=re.MULTILINE)
    return [Document(section.strip(), metadata={"source": path.name,
        "section": section.splitlines()[0], "document_type": "analysis_summary"})
        for section in sections if section.strip()]


def _csv_documents() -> list[Document]:
    documents = []
    for row in read_csv("error_analysis_per_image.csv"):
        model, tier = row.get("Model", ""), row.get("Tier", "")
        character, pattern = row.get("Character_or_Setting", ""), row.get("Primary_Error_Pattern", "")
        content = (f"{model} 错误案例｜{tier}｜{character}\n主要错误：{pattern}\n"
                   f"标签：{row.get('Death_Tag', '')}\n详情：{row.get('Detailed_Error_Description', '')}")
        metadata = {"source": "error_analysis_per_image.csv", "document_type": "error_case",
                    "model": model, "tier": tier, "index": int(row["Index"])}
        if character:
            metadata["character"] = character
        if pattern:
            metadata["error_type"] = pattern
        if tier.startswith("Tier "):
            metadata["task"] = f"Task {tier[-1]}"
        documents.append(Document(content, metadata=metadata))
    for row in read_csv("error_pattern_summary.csv"):
        model, tier, pattern = row.get("Model", ""), row.get("Tier", ""), row.get("Error_Pattern", "")
        content = (f"{model}｜{tier} 错误类型统计\n错误类型：{pattern}\n数量：{row.get('Count', '')}\n"
                   f"占该模型错误比例：{row.get('Percentage_of_Model_Errors', '')}")
        metadata = {"source": "error_pattern_summary.csv", "document_type": "pattern_summary",
                    "model": model, "tier": tier, "error_type": pattern}
        if tier.startswith("Tier "):
            metadata["task"] = f"Task {tier[-1]}"
        documents.append(Document(content, metadata=metadata))
    for filename in ("qwen_doubao_task_accuracy.csv", "task_model_accuracy_4x4.csv"):
        for row in read_csv(filename):
            model, tier = row.get("Model", "").split()[-1], row.get("Tier", "")
            metadata = {"source": filename, "document_type": "accuracy_summary", "model": model, "tier": tier}
            if tier.startswith("Tier "):
                metadata["task"] = f"Task {tier[-1]}"
            documents.append(Document("｜".join(f"{key}：{value}" for key, value in row.items()), metadata=metadata))
    return documents


def load_eval_documents() -> list[Document]:
    global _DOC_CACHE, _DOC_CACHE_SIGNATURE
    signature = data_signature()
    if _DOC_CACHE is None or signature != _DOC_CACHE_SIGNATURE:
        _DOC_CACHE = _workbook_documents() + _markdown_documents() + _csv_documents()
        _DOC_CACHE_SIGNATURE = signature
        get_vector_store.cache_clear()
        from retrieval import clear_retrieval_cache
        clear_retrieval_cache()
    return _DOC_CACHE


def index_version(documents: list[Document], chunk_size: int) -> str:
    digest = hashlib.sha256()
    for document in documents:
        digest.update(document.page_content.encode("utf-8"))
        digest.update(json.dumps(document.metadata, ensure_ascii=False, sort_keys=True).encode("utf-8"))
    digest.update(f"{EMBEDDING_MODEL}|{chunk_size}|{min(80, chunk_size // 5)}".encode())
    return digest.hexdigest()[:16]


@lru_cache(maxsize=8)
def _cached_vector_store(chunk_size: int, version: str) -> Chroma:
    persist_directory = DB_DIR / "bge-small-zh-v1.5" / f"chunk_{chunk_size}" / version
    embeddings = ChineseSemanticEmbeddings()
    splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=min(80, chunk_size // 5))
    documents = splitter.split_documents(load_eval_documents())
    store = Chroma(collection_name=f"eval_rules_bge_zh_{chunk_size}", embedding_function=embeddings,
        persist_directory=str(persist_directory), collection_metadata={"hnsw:space": "cosine"})
    for doc in documents:
        doc.metadata["index_version"] = version
    ids = [hashlib.sha1(f"{EMBEDDING_MODEL}|{chunk_size}|{i}|{doc.page_content}|"
        f"{json.dumps(doc.metadata, ensure_ascii=False, sort_keys=True)}".encode("utf-8")).hexdigest()
        for i, doc in enumerate(documents)]
    if set(store.get(include=[])["ids"]) != set(ids):
        store.reset_collection()
        store.add_documents(documents, ids=ids)
    return store


def get_vector_store(chunk_size: int = 500) -> Chroma:
    return _cached_vector_store(chunk_size, index_version(load_eval_documents(), chunk_size))


get_vector_store.cache_clear = _cached_vector_store.cache_clear
